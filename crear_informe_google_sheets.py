from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\jhonb\.codex\visualizations\2026\10\12\01a1314b-5f13-7051-b543-fd8df12e5ad5\Informe procedimiento Google Sheets.docx"


def set_font(run, size=11, bold=False, color="000000"):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_el = tc_pr.first_child_found_in("w:tcBorders")
    if borders_el is None:
        borders_el = OxmlElement("w:tcBorders")
        tc_pr.append(borders_el)
    for edge in ("top", "left", "bottom", "right"):
        el = borders_el.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders_el.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def add_text(doc, text, before=0, after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text))
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_font(r, size=14, bold=True)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    set_font(p.add_run(text))
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(7)
set_font(title.add_run("Procedimiento de vinculacion con Google Sheets"), size=19, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(20)
set_font(subtitle.add_run("Sistema de registro de citas del Centro de Salud San Juan de Dios Pisco"), size=11, color="404040")

add_heading(doc, "Introduccion")
add_text(doc, "En este trabajo implementé la vinculación entre el formulario de reserva de citas de la página web del Centro de Salud San Juan de Dios Pisco y una hoja de cálculo de Google Sheets. El objetivo fue que cada solicitud registrada desde el portal del paciente quede almacenada automáticamente en una base de datos sencilla, ordenada y accesible para el personal encargado.")
add_text(doc, "La solución se desarrolló con Google Apps Script como intermediario entre la página web y Google Sheets. De esta forma, no fue necesario instalar un servidor propio para registrar inicialmente las citas.")

add_heading(doc, "Objetivo")
add_text(doc, "Mi objetivo fue registrar las solicitudes de cita generadas desde el portal del paciente en una hoja de cálculo de Google Sheets, conservando la información principal de cada solicitud y su estado.")

add_heading(doc, "Herramientas utilizadas")
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["Herramienta", "Uso dentro del proyecto"]
for cell, text in zip(table.rows[0].cells, headers):
    shade(cell, "0B6F85")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), size=10, bold=True, color="FFFFFF")
    borders(cell)
for left, right in [
    ("Página web HTML CSS JavaScript", "Formulario de reserva de citas y envío de datos."),
    ("Google Sheets", "Almacenamiento y visualización de las solicitudes recibidas."),
    ("Google Apps Script", "Servicio web que recibe los datos y los registra en la hoja."),
    ("GitHub", "Respaldo y publicación de los cambios realizados en el proyecto."),
]:
    cells = table.add_row().cells
    for idx, text in enumerate((left, right)):
        cells[idx].text = ""
        cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.space_before = Pt(3)
        set_font(p.add_run(text), size=10)
        borders(cells[idx])

add_heading(doc, "Procedimiento realizado")
steps = [
    "Creé una hoja de cálculo de Google Sheets para registrar las citas. Dentro de ella preparé una pestaña llamada Citas con los campos: fecha de registro, código de solicitud, paciente, DNI, especialidad, profesional, fecha solicitada, hora solicitada y estado.",
    "Abrí el editor de Google Apps Script desde la hoja de cálculo y programé una función para preparar la pestaña y sus encabezados. También configuré una función doPost para recibir los datos enviados desde la página web.",
    "En el código de Apps Script programé el registro de una nueva fila cada vez que se recibe una solicitud. Al llegar los datos, el sistema guarda la fecha, el DNI del paciente, la especialidad elegida, el profesional, la fecha y hora solicitadas, además del estado inicial Solicitud registrada.",
    "Implementé acciones de reprogramación y cancelación. Para evitar que se modifique una cita equivocada, cada solicitud utiliza un código único. Cuando una cita se cancela o se reprograma, el sistema actualiza la fila existente y cambia su estado, sin crear registros duplicados.",
    "Publiqué el proyecto de Apps Script como una aplicación web. Luego obtuve la dirección web de ejecución generada por Google y la agregué en la configuración de JavaScript de la página.",
    "Modifiqué el formulario de reserva de citas del portal del paciente. Al confirmar una solicitud, el formulario reúne los datos ingresados, genera un código de solicitud y los envía en formato JSON al servicio de Google Apps Script.",
    "Realicé una prueba registrando una cita desde el portal. Finalmente comprobé que los datos aparecieran correctamente en Google Sheets y que el estado se actualizara al cancelar o reprogramar una cita.",
]
for step in steps:
    add_numbered(doc, step)

add_heading(doc, "Resultado obtenido")
add_text(doc, "Como resultado, la página web ahora registra las solicitudes de citas en Google Sheets. Esto permite contar con un listado centralizado de citas y facilita su revisión desde una cuenta de Google autorizada. Además, el portal muestra al paciente la cita que acaba de registrar y permite solicitar su cancelación o reprogramación.")

add_heading(doc, "Consideraciones finales")
add_text(doc, "La integración desarrollada funciona como una solución práctica para el proyecto demostrativo. Para utilizarla en un entorno real de atención médica, recomendaría incorporar autenticación de usuarios, validación adicional de la información, permisos restringidos para la hoja de cálculo y medidas de protección para los datos personales de los pacientes.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Centro de Salud San Juan de Dios Pisco  Informe de proyecto"), size=8, color="666666")

doc.core_properties.title = "Procedimiento de vinculacion con Google Sheets"
doc.core_properties.author = ""
doc.save(OUT)
